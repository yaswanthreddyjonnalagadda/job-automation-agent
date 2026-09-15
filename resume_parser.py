"""
Resume parsing: extract raw text from PDF/DOCX and (optionally) turn it into
structured data via Claude.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import docx
from pypdf import PdfReader

logger = logging.getLogger(__name__)


@dataclass
class ResumeData:
    source_path: str
    raw_text: str
    structured: dict[str, Any] = field(default_factory=dict)


def _extract_text_from_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages):
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            logger.warning("Failed to extract text from PDF page %d of %s", i, path)
    return "\n".join(pages).strip()


def _extract_text_from_docx(path: Path) -> str:
    document = docx.Document(str(path))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)
    return "\n".join(paragraphs).strip()


def extract_text(path: str | Path) -> str:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Resume file not found: {path}")

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = _extract_text_from_pdf(path)
    elif suffix == ".docx":
        text = _extract_text_from_docx(path)
    else:
        raise ValueError(f"Unsupported resume format: {suffix} (use .pdf or .docx)")

    if not text:
        raise ValueError(f"No extractable text found in resume: {path}")

    logger.info("Extracted %d characters of resume text from %s", len(text), path)
    return text


def parse_resume(path: str | Path) -> ResumeData:
    """Extract raw text only. Call claude_integration.structure_resume() separately
    to get a structured (skills/experience/education) breakdown via Claude."""
    path = Path(path)
    raw_text = extract_text(path)
    return ResumeData(source_path=str(path), raw_text=raw_text)
